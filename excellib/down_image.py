"""
从Excel指定列下载图片并插入到相邻列

功能特性：
- 支持替换原单元格或插入到新列
- 自动调整行高和列宽以适应图片
- 图片固定在单元格（随单元格移动）
- 防止内容溢出（文本自动换行+居中对齐）
- 可自定义图片尺寸
- 支持批量处理

图片固定与防溢出：
- 使用anchor锚点固定图片位置
- 图片会跟随单元格移动（插入/删除行列时）
- 调整行高列宽使图片完美适配单元格
- 设置单元格对齐方式防止文本溢出

单位换算说明：
┌──────────────┬─────────────┬──────────────┐
│ Excel单位    │ 用途        │ 转换系数     │
├──────────────┼─────────────┼──────────────┤
│ 点(point)    │ 行高        │ 像素×0.75    │
│ 字符宽度     │ 列宽        │ 像素×0.14    │
│ 像素(pixel)  │ 图片尺寸    │ 基准单位     │
└──────────────┴─────────────┴──────────────┘

示例：150×150像素图片 → 行高112.5点 + 列宽21字符
"""
import os
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment
import requests
from pathlib import Path


def download_images_from_column(
    excel_path: str,
    url_column: str,
    image_column: str = None,
    sheet_name: str = None,
    start_row: int = 2,
    image_width: int = 100,
    image_height: int = 100,
    output_path: str = None,
    replace_url: bool = True,
    auto_row_height: bool = True,
    row_height_ratio: float = 0.75,
    auto_col_width: bool = True,
    col_width_ratio: float = 0.14,
    wrap_text: bool = True,
    align_center: bool = True
):
    """
    从Excel指定列下载图片并插入到目标列

    参数:
        excel_path: Excel文件路径
        url_column: URL所在列名（如'A'）或列索引（如1）
        image_column: 图片插入的列名（如'B'），默认替换原单元格
        sheet_name: 工作表名称，默认使用活动工作表
        start_row: 开始行号（默认2，跳过表头）
        image_width: 图片宽度（像素）
        image_height: 图片高度（像素）
        output_path: 输出文件路径，默认覆盖原文件
        replace_url: 是否替换原URL单元格（默认True）
        auto_row_height: 是否自动调整行高（默认True）
        row_height_ratio: 行高调整系数（默认0.75，像素转点的比率）
        auto_col_width: 是否自动调整列宽（默认True）
        col_width_ratio: 列宽调整系数（默认0.14，像素转字符宽度）
        wrap_text: 是否设置单元格文本自动换行，防止溢出（默认True）
        align_center: 是否设置单元格内容居中对齐（默认True）
    """
    # 加载工作簿
    wb = load_workbook(excel_path)
    ws = wb[sheet_name] if sheet_name else wb.active

    # 处理列名
    if isinstance(url_column, str) and url_column.isalpha():
        url_col_idx = ord(url_column.upper()) - ord('A') + 1
    else:
        url_col_idx = int(url_column)

    if image_column is None:
        # 默认替换原单元格
        img_col_idx = url_col_idx
    elif isinstance(image_column, str) and image_column.isalpha():
        img_col_idx = ord(image_column.upper()) - ord('A') + 1
    else:
        img_col_idx = int(image_column)

    # 获取列字母
    img_col_letter = chr(ord('A') + img_col_idx - 1)

    # 获取总行数
    max_row = ws.max_row
    total_rows = max_row - start_row + 1

    replace_mode = (img_col_idx == url_col_idx)
    mode_text = "替换原单元格" if replace_mode else f"插入到{img_col_letter}列"
    print(f"工作表总行数: {max_row}, 待处理行数: {total_rows}")
    print(f"处理模式: {mode_text}")
    print(f"图片尺寸: {image_width}x{image_height} 像素")
    if auto_row_height:
        calculated_height = image_height * row_height_ratio
        print(f"自动调整行高: 启用 (行高={calculated_height:.1f}点)")
    if auto_col_width:
        calculated_width = image_width * col_width_ratio
        print(f"自动调整列宽: 启用 (列宽={calculated_width:.1f}字符)")

    # 遍历行
    success_count = 0
    error_count = 0
    skip_count = 0

    for row in range(start_row, max_row + 1):
        cell = ws.cell(row=row, column=url_col_idx)
        url = cell.value

        # 如果单元格为空，跳过
        if not url or not str(url).strip():
            skip_count += 1
            continue

        try:
            current_num = row - start_row + 1
            print(f"正在下载第 {row} 行图片 ({current_num}/{total_rows}): {url}")

            # 下载图片
            response = requests.get(url, timeout=10)
            response.raise_for_status()

            # 创建图片对象
            img_data = BytesIO(response.content)
            img = Image(img_data)

            # 设置图片大小
            img.width = image_width
            img.height = image_height

            # 插入图片到单元格并固定
            # anchor设置图片左上角位置
            img.anchor = f"{img_col_letter}{row}"
            
            # 设置图片锁定模式（默认oneCell：固定位置但不随单元格调整大小）
            # 如果要让图片随单元格移动和调整，可以在图片对象添加后修改属性
            ws.add_image(img)

            # 如果是替换模式且图片列与URL列相同，清空URL内容
            if replace_url and img_col_idx == url_col_idx:
                cell.value = None

            # 设置单元格对齐方式，防止内容溢出
            if wrap_text or align_center:
                alignment = Alignment(
                    horizontal='center' if align_center else None,
                    vertical='center' if align_center else None,
                    wrap_text=wrap_text
                )
                cell.alignment = alignment

            # 调整行高以适应图片
            if auto_row_height:
                ws.row_dimensions[row].height = image_height * row_height_ratio

            success_count += 1

        except Exception as e:
            print(f"第 {row} 行图片下载失败: {str(e)}")
            error_count += 1

    # 调整列宽（固定图片在单元格中的关键设置）
    if auto_col_width:
        ws.column_dimensions[img_col_letter].width = image_width * col_width_ratio

    # 保存文件
    save_path = output_path if output_path else excel_path
    wb.save(save_path)

    print(f"\n处理完成！")
    print(f"成功: {success_count} 张")
    print(f"失败: {error_count} 张")
    print(f"跳过: {skip_count} 行")
    print(f"保存路径: {save_path}")


if __name__ == "__main__":
    # 示例用法
    import sys

    # if len(sys.argv) < 3:
    #     print("用法: python down_image.py <Excel文件路径> <URL列名或列号>")
    #     print("示例: python down_image.py data.xlsx A")
    #     print("示例: python down_image.py data.xlsx 1")
    #     sys.exit(1)

    # excel_file = sys.argv[1]
    # url_col = sys.argv[2]

    # # 可选参数
    # sheet = sys.argv[3] if len(sys.argv) > 3 else None

    excel_file = "a.xlsx"
    url_col = "D"
    out_file = "b.xlsx"

    download_images_from_column(
        excel_path=excel_file,
        url_column=url_col,
        output_path=out_file,
        image_width=150,          # 图片宽度（像素）
        image_height=150,         # 图片高度（像素）
        auto_row_height=True,     # 自动调整行高
        row_height_ratio=0.75,    # 行高系数：像素→点
        auto_col_width=True,      # 自动调整列宽
        col_width_ratio=0.14,     # 列宽系数：像素→字符
        wrap_text=True,           # 文本自动换行，防止溢出
        align_center=True         # 内容居中对齐
    )
